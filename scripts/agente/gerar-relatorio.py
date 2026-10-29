#!/usr/bin/env python3
"""Gera o relatório do Agente Licitacon em Word (.docx) a partir do resultado de uma análise (JSON).

- O texto literal de cada dispositivo citado vem da base oficial (agente/normas/<norma>.json).
  Se a norma ainda não estiver na base, o relatório marca "[texto a conferir na fonte oficial]".
- Fundamento, gravidade, risco e dica de cada item vêm do roteiro (agente/roteiros/<regime>.json);
  a análise só precisa informar o resultado, o trecho do documento e a explicação do caso concreto.
- A jurisprudência vem exclusivamente da Biblioteca (acervo.json), com link para a fonte oficial.

Uso: python3 scripts/agente/gerar-relatorio.py analise.json saida.docx
Estrutura do JSON de análise: ver agente/relatorio/estrutura.md
"""
import datetime, json, pathlib, re, sys, unicodedata
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

RAIZ = pathlib.Path(__file__).resolve().parents[2]
COR = {"alta": "8A1230", "media": "8A4A12", "baixa": "0F6B7E", "ok": "0A7A50", "titulo": "0A1F2A", "fraco": "6A7B82"}
FUNDO = {"alta": "FCEEF1", "media": "FDF1E3", "baixa": "E9F3F5"}
GRAV = {"alta": "Alta", "media": "Média", "baixa": "Baixa"}
RESULT = {"nao_conforme": "Não conforme", "conforme": "Conforme", "parcial": "Parcialmente conforme",
          "nao_aplicavel": "Não se aplica", "nao_verificavel": "Não verificável (documento ausente ou ilegível)"}


# ------------------------------------------------------------------ base normativa e Biblioteca
_normas = {}
def norma(nid):
    if nid not in _normas:
        p = RAIZ / "agente" / "normas" / f"{nid}.json"
        _normas[nid] = json.loads(p.read_text(encoding="utf-8")) if p.exists() else None
    return _normas[nid]

NOME_CURTO = {"lei-14133": "Lei 14.133/2021", "lei-13303": "Lei 13.303/2016", "lei-13019": "Lei 13.019/2014",
              "lei-10192": "Lei 10.192/2001", "lei-4320": "Lei 4.320/1964"}

def rotulo(f):
    n = norma(f["norma"])
    disp = f.get("disp", "caput")
    rot = None
    if n:
        art = next((a for a in n["artigos"] if a["num"] == f["art"]), None)
        d = next((x for x in (art or {}).get("dispositivos", []) if x["id"] == disp), None)
        rot = d["rotulo"] if d else None
    if not rot:  # monta o rótulo a partir do id (p1.I.a → § 1º, inciso I, alínea a)
        partes = []
        for p in disp.split("."):
            if p == "caput": partes.append("caput")
            elif p == "pu": partes.append("parágrafo único")
            elif p.startswith("p") and p[1:2].isdigit(): partes.append("§ " + p[1:] + ("º" if p[1:].isdigit() and int(p[1:]) < 10 else ""))
            elif re.fullmatch(r"[IVXLCDM]+(-[A-Z])?", p): partes.append("inciso " + p)
            elif re.fullmatch(r"[a-z]", p): partes.append("alínea " + p)
            else: partes.append("item " + p)
        rot = ", ".join(partes)
    rot = "" if rot == "caput" else ", " + rot
    return f"art. {f['art']}{rot}, da {NOME_CURTO.get(f['norma'], f['norma'])}"

def texto_legal(f):
    n = norma(f["norma"])
    if not n:
        return None
    art = next((a for a in n["artigos"] if a["num"] == f["art"]), None)
    if not art:
        return None
    disp = f.get("disp", "caput")
    d = next((x for x in art["dispositivos"] if x["id"] == disp), None)
    if not d:
        return None
    t = d["texto"]
    # inciso de um caput/parágrafo: mostra também o enunciado ao qual pertence
    if "." in disp or re.fullmatch(r"[IVXLCDM]+(-[A-Z])?", disp):
        pai_id = disp.rsplit(".", 1)[0] if "." in disp else "caput"
        pai = next((x for x in art["dispositivos"] if x["id"] == pai_id), None)
        if pai and pai["id"] != disp:
            t = pai["texto"].rstrip(":") + ": […] " + t
    return t

_acervo = None
def _norm(s):
    return unicodedata.normalize("NFD", (s or "").lower()).encode("ascii", "ignore").decode()

def jurisprudencia(temas, n=3):
    """Busca simples na Biblioteca: pontua pela presença dos termos do tema em título, tese, tags e tema."""
    global _acervo
    if _acervo is None:
        p = RAIZ / "acervo.json"
        _acervo = json.loads(p.read_text(encoding="utf-8")).get("itens", []) if p.exists() else []
    termos = [_norm(t) for t in temas or [] if t]
    res = []
    for it in _acervo:
        campo = _norm(" ".join([it.get("titulo", ""), it.get("tese", ""), " ".join(it.get("tags") or []), it.get("tema", "")]))
        achados = [t for t in termos if t in campo]
        s = sum((3 if len(t.split()) > 1 else 1) for t in achados)
        contexto = any(k in campo for k in ("licita", "contrat", "edital", "pregao", "registro de precos", "termo de referencia"))
        # exige contexto de licitação/contrato e mais de um termo (ou uma expressão composta)
        if contexto and s >= 2:
            res.append((s, it.get("data", ""), it))
    res.sort(key=lambda x: (x[0], x[1]), reverse=True)
    return [r[2] for r in res[:n]]


# ------------------------------------------------------------------ formatação
def sombrear(celula, cor):
    tcpr = celula._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), cor)
    tcpr.append(shd)

def par(doc, texto="", negrito=False, cor=None, tam=None, italico=False, alinhamento=None, depois=6):
    p = doc.add_paragraph()
    if texto:
        r = p.add_run(texto); r.bold = negrito; r.italic = italico
        if cor: r.font.color.rgb = RGBColor.from_string(cor)
        if tam: r.font.size = Pt(tam)
    if alinhamento: p.alignment = alinhamento
    p.paragraph_format.space_after = Pt(depois)
    return p

def rotulo_valor(doc, rot, valor, cor=None):
    p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
    a = p.add_run(rot + ": "); a.bold = True; a.font.color.rgb = RGBColor.from_string(COR["titulo"])
    b = p.add_run(valor)
    if cor: b.font.color.rgb = RGBColor.from_string(cor)
    return p

def link(paragrafo, texto, url):
    rid = paragrafo.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), rid)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0F6B7E"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = texto; r.append(t); h.append(r)
    paragrafo._p.append(h)


# ------------------------------------------------------------------ relatório
def gerar(analise, saida):
    roteiro = json.loads((RAIZ / "agente" / "roteiros" / f"{analise['regime']}.json").read_text(encoding="utf-8"))
    itens = {i["id"]: i for i in roteiro["itens"]}
    docs_nome = roteiro["documentos"]
    resultados = analise["resultados"]
    for r in resultados:
        if r["item"] not in itens:
            raise SystemExit(f"Item desconhecido no roteiro: {r['item']}")

    doc = Document()
    st = doc.styles["Normal"]; st.font.name = "Calibri"; st.font.size = Pt(10.5)
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.2); s.top_margin = s.bottom_margin = Cm(2)

    # cabeçalho
    par(doc, "LICITACON · AGENTE DE ANÁLISE", True, COR["ok"], 9, depois=2)
    par(doc, "Relatório de análise da fase preparatória", True, COR["titulo"], 20, depois=4)
    par(doc, analise["processo"].get("objeto", ""), False, COR["fraco"], 12, depois=12)
    for rot, chave in [("Órgão", "orgao"), ("Processo", "numero"), ("Modalidade / critério", "modalidade")]:
        if analise["processo"].get(chave):
            rotulo_valor(doc, rot, analise["processo"][chave])
    rotulo_valor(doc, "Regime jurídico identificado", analise.get("regime_texto", NOME_CURTO.get(analise["regime"], analise["regime"])))
    rotulo_valor(doc, "Documentos analisados", "; ".join(f"{docs_nome.get(d['tipo'], d['tipo'])} ({d.get('arquivo', '')}{', ' + str(d['paginas']) + ' p.' if d.get('paginas') else ''})" for d in analise["documentos"]))
    rotulo_valor(doc, "Data da análise", analise.get("data", datetime.date.today().strftime("%d/%m/%Y")))
    rotulo_valor(doc, "Versão do roteiro", f"{roteiro['regime']} · {roteiro['versao']}")

    aviso = doc.add_paragraph(); aviso.paragraph_format.space_before = Pt(10)
    r = aviso.add_run("Natureza deste relatório: análise técnica automatizada de apoio à instrução do processo, em gestão de licitações e contratos. "
                      "Não é parecer jurídico e não substitui a análise jurídica prévia obrigatória (art. 53 da Lei 14.133/2021) nem a decisão da autoridade competente. "
                      "Confira cada apontamento no documento e na fonte oficial antes de adotá-lo.")
    r.italic = True; r.font.size = Pt(9); r.font.color.rgb = RGBColor.from_string(COR["fraco"])

    # resumo executivo
    falhas = [r for r in resultados if r["resultado"] in ("nao_conforme", "parcial")]
    cont = {g: sum(1 for r in falhas if itens[r["item"]]["gravidade"] == g) for g in ("alta", "media", "baixa")}
    doc.add_heading("1. Resumo executivo", level=1)
    sit = "Há falhas graves que recomendam correção." if cont["alta"] else (
          "Há ajustes recomendados, sem falha grave identificada." if falhas else "Nenhuma não conformidade identificada nos itens verificados.")
    par(doc, sit, True, COR["alta"] if cont["alta"] else (COR["media"] if falhas else COR["ok"]), 12)
    t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.style = "Table Grid"
    for i, (rot, val, g) in enumerate([("Gravidade alta", cont["alta"], "alta"), ("Gravidade média", cont["media"], "media"),
                                       ("Gravidade baixa", cont["baixa"], "baixa"), ("Itens verificados", len(resultados), None)]):
        c = t.rows[0].cells[i]; c.text = ""
        p = c.paragraphs[0]; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a = p.add_run(str(val) + "\n"); a.bold = True; a.font.size = Pt(18)
        if g: a.font.color.rgb = RGBColor.from_string(COR[g]); sombrear(c, FUNDO[g])
        b = p.add_run(rot); b.font.size = Pt(9)
    if analise.get("resumo"):
        par(doc, ""); par(doc, analise["resumo"])

    # achados
    doc.add_heading("2. Falhas identificadas", level=1)
    if not falhas:
        par(doc, "Nenhuma falha identificada nos itens aplicáveis.", italico=True)
    ordem = {"alta": 0, "media": 1, "baixa": 2}
    for n, r in enumerate(sorted(falhas, key=lambda x: ordem[itens[x["item"]]["gravidade"]]), 1):
        it = itens[r["item"]]; g = it["gravidade"]
        h = doc.add_heading(level=2)
        a = h.add_run(f"2.{n}. {it['titulo']}  "); a.font.color.rgb = RGBColor.from_string(COR["titulo"])
        b = h.add_run(f"[{GRAV[g]} · {RESULT[r['resultado']]}]"); b.font.color.rgb = RGBColor.from_string(COR[g]); b.font.size = Pt(10)
        rotulo_valor(doc, "Documento", f"{docs_nome.get(it['doc'], it['doc'])}" + (f", {r['local']}" if r.get("local") else ""))
        if r.get("trecho"):
            p = par(doc, "“" + r["trecho"] + "”", italico=True, cor=COR["fraco"]); p.paragraph_format.left_indent = Cm(0.8)
        rotulo_valor(doc, "O que está em desacordo", r["constatacao"])
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(3)
        x = p.add_run("Fundamento legal: "); x.bold = True; x.font.color.rgb = RGBColor.from_string(COR["titulo"])
        p.add_run("; ".join(rotulo(f) for f in it["fundamento"]) + ".")
        for f in it["fundamento"]:
            tl = texto_legal(f)
            q = par(doc, (rotulo(f)[0].upper() + rotulo(f)[1:] + ": “" + tl + "”") if tl else (rotulo(f) + ": [texto a conferir na fonte oficial]"),
                    cor=COR["fraco"], tam=9, italico=True, depois=3)
            q.paragraph_format.left_indent = Cm(0.8)
        rotulo_valor(doc, "Por que está errado", r.get("porque") or it["verificar"])
        rotulo_valor(doc, "Risco de manter", r.get("risco") or it["risco"], COR[g])
        rotulo_valor(doc, "Como corrigir", r.get("correcao") or it["dica"])
        if r.get("redacao_sugerida"):
            p = par(doc, "Redação sugerida: " + r["redacao_sugerida"], cor=COR["ok"]); p.paragraph_format.left_indent = Cm(0.8)
        if "jurisprudencia_ids" in r:  # escolhida pela análise (o agente seleciona entre os candidatos da Biblioteca)
            jurisprudencia([])  # garante a Biblioteca carregada
            por_id = {j.get("id"): j for j in _acervo}
            js = [por_id[i] for i in r["jurisprudencia_ids"] if i in por_id]
        else:
            js = jurisprudencia(r.get("temas") or it.get("temas"))
        p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
        x = p.add_run("Jurisprudência na Biblioteca Licitacon: "); x.bold = True; x.font.color.rgb = RGBColor.from_string(COR["titulo"])
        if not js:
            p.add_run("nenhuma decisão sobre o tema localizada na Biblioteca até esta data.")
        for j in js:
            q = doc.add_paragraph(style="List Bullet"); q.paragraph_format.space_after = Pt(2)
            q.add_run(f"{j.get('orgao', '')} · {j.get('numero', '')} ({j.get('data', '')[:10]}): ").bold = True
            q.add_run((j.get("titulo") or "") + " ")
            if j.get("fonte_url"):
                link(q, "abrir na fonte oficial", j["fonte_url"])

    # dicas
    dicas = [r for r in resultados if r["resultado"] == "conforme" and r.get("dica")]
    doc.add_heading("3. Recomendações de melhoria", level=1)
    par(doc, "Pontos sem irregularidade, mas que podem ser aperfeiçoados:", italico=True, cor=COR["fraco"])
    if not dicas:
        par(doc, "Nenhuma recomendação adicional.", italico=True)
    for r in dicas:
        it = itens[r["item"]]
        q = doc.add_paragraph(style="List Bullet")
        q.add_run(it["titulo"] + ": ").bold = True
        q.add_run(r["dica"] + " (" + "; ".join(rotulo(f) for f in it["fundamento"]) + ")")

    # lista de verificação completa
    doc.add_heading("4. Lista de verificação", level=1)
    par(doc, "Todos os itens do roteiro aplicados a este processo.", italico=True, cor=COR["fraco"])
    t = doc.add_table(rows=1, cols=4); t.style = "Table Grid"
    for i, h in enumerate(["Item", "Documento", "Resultado", "Fundamento"]):
        c = t.rows[0].cells[i]; c.text = h; c.paragraphs[0].runs[0].bold = True; sombrear(c, "EAF2F1")
    for r in resultados:
        it = itens[r["item"]]; row = t.add_row().cells
        row[0].text = it["titulo"]; row[1].text = docs_nome.get(it["doc"], it["doc"])
        row[2].text = RESULT[r["resultado"]]
        row[3].text = "; ".join(rotulo(f) for f in it["fundamento"])
        if r["resultado"] in ("nao_conforme", "parcial"):
            sombrear(row[2], FUNDO[it["gravidade"]])
    for row in t.rows:
        row.cells[0].width = Cm(5.5); row.cells[1].width = Cm(3.2); row.cells[2].width = Cm(3.2); row.cells[3].width = Cm(5)
        for c in row.cells:
            for p in c.paragraphs:
                for rr in p.runs: rr.font.size = Pt(8.5)

    # metodologia
    doc.add_heading("5. Metodologia e limites", level=1)
    for txt in [
        "Cada item do roteiro foi verificado nos documentos enviados; o trecho citado indica onde a constatação foi feita.",
        "Os dispositivos legais citados reproduzem o texto oficial vigente na base do Licitacon na data da análise.",
        "A jurisprudência indicada vem exclusivamente da Biblioteca Licitacon, com link para a fonte oficial; a ausência de decisão listada não significa ausência de jurisprudência sobre o tema.",
        "Cálculos de planilha são conferidos por programa, a partir do arquivo em Excel; planilhas em PDF não permitem conferir fórmulas.",
        "Normas estaduais e municipais, regulamentos internos e orientações locais só são considerados quando enviados com os documentos.",
    ]:
        doc.add_paragraph(txt, style="List Bullet")

    doc.save(saida)
    print("relatório gravado:", saida)


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    gerar(json.loads(pathlib.Path(sys.argv[1]).read_text(encoding="utf-8")), sys.argv[2])
