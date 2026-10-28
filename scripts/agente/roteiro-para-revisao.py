#!/usr/bin/env python3
"""Gera um Word com os roteiros de verificação, para revisão do especialista.
Uso: python3 scripts/agente/roteiro-para-revisao.py agente/roteiros/lei-14133.json saida.docx"""
import json, pathlib, sys
from docx import Document
from docx.shared import Cm, Pt, RGBColor

r = json.loads(pathlib.Path(sys.argv[1]).read_text(encoding="utf-8"))
GRAV = {"alta": ("Alta", "8A1230"), "media": ("Média", "8A4A12"), "baixa": ("Baixa", "0F6B7E")}
OBR = {"sim": "Obrigatório", "condicional": "Condicional", "recomendado": "Recomendado"}
NOME = {"lei-14133": "Lei 14.133/2021", "lei-10192": "Lei 10.192/2001", "lei-13303": "Lei 13.303/2016", "lei-13019": "Lei 13.019/2014", "lei-4320": "Lei 4.320/1964"}

def disp(d):
    out = []
    for p in d.split("."):
        if p == "caput": continue
        if p == "pu": out.append("parágrafo único")
        elif p.startswith("p") and p[1:].replace("-", "").isalnum() and p[1:2].isdigit(): out.append("§ " + p[1:] + ("º" if p[1:].isdigit() and int(p[1:]) < 10 else ""))
        elif len(p) == 1 and p.islower(): out.append("alínea " + p)
        else: out.append("inc. " + p)
    return (", " + ", ".join(out)) if out else ""

doc = Document()
doc.styles["Normal"].font.name = "Calibri"; doc.styles["Normal"].font.size = Pt(10)
for s in doc.sections: s.left_margin = s.right_margin = Cm(2)
doc.add_heading("Agente Licitacon — roteiros de verificação (" + NOME.get(r["regime"], r["regime"]) + ")", 0)
doc.add_paragraph("Versão " + r["versao"] + ". Para revisão: corrija fundamentos, gravidades e textos diretamente neste arquivo (de preferência com o controle de alterações ligado) e devolva. Cada item vira uma verificação automática do agente.")
for chave, nome in r["documentos"].items():
    itens = [i for i in r["itens"] if i["doc"] == chave]
    if not itens: continue
    doc.add_heading(nome, 1)
    for i in itens:
        g, cor = GRAV[i["gravidade"]]
        h = doc.add_heading(level=2); h.add_run(i["id"] + " · " + i["titulo"] + "  ")
        x = h.add_run("[" + g + " · " + OBR[i["obrigatorio"]] + "]"); x.font.color.rgb = RGBColor.from_string(cor); x.font.size = Pt(10)
        for rot, val in [("Verificar", i["verificar"]),
                         ("Quando se aplica", i.get("quando")),
                         ("Fundamento", "; ".join("art. " + f["art"] + disp(f["disp"]) + " da " + NOME.get(f["norma"], f["norma"]) for f in i["fundamento"])),
                         ("Risco de manter", i["risco"]), ("Dica de ajuste", i["dica"]),
                         ("Termos para jurisprudência", ", ".join(i.get("temas", [])))]:
            if not val: continue
            p = doc.add_paragraph(); p.paragraph_format.space_after = Pt(2)
            p.add_run(rot + ": ").bold = True; p.add_run(val)
doc.save(sys.argv[2]); print("gravado", sys.argv[2])
