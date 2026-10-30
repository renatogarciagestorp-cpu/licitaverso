#!/usr/bin/env python3
"""Converte planilha Excel (.xlsx/.xlsm; .xls/.ods convertidos pelo LibreOffice) em texto para o agente.
Cada célula preenchida vira "COLUNA=valor"; quando a célula tem fórmula, ela vem entre colchetes depois do valor calculado,
para o agente conferir como o BDI, os encargos e os totais foram aplicados.

Uso: python3 scripts/agente/extrair-planilha.py planilha.xlsx saida.txt [--max 300000]
"""
import datetime, pathlib, subprocess, sys, tempfile

import openpyxl
from openpyxl.utils import get_column_letter

MAX = 300_000


def converter(caminho):
    """.xls / .ods → .xlsx (LibreOffice)."""
    tmp = pathlib.Path(tempfile.mkdtemp())
    subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp), str(caminho)],
                   capture_output=True, timeout=300)
    saida = next(tmp.glob("*.xlsx"), None)
    if not saida:
        sys.exit("Não foi possível converter a planilha para .xlsx.")
    return saida


def fmt(v, formato):
    if v is None:
        return ""
    if isinstance(v, bool):
        return "VERDADEIRO" if v else "FALSO"
    if isinstance(v, (datetime.datetime, datetime.date)):
        return v.strftime("%d/%m/%Y")
    if isinstance(v, float):
        if "%" in (formato or ""):
            return f"{v * 100:.4g}%"
        return f"{v:.6f}".rstrip("0").rstrip(".")
    if isinstance(v, int) and "%" in (formato or ""):
        return f"{v * 100}%"
    return " ".join(str(v).split())


def extrair(caminho, limite=MAX):
    p = pathlib.Path(caminho)
    if p.suffix.lower() not in (".xlsx", ".xlsm"):
        p = converter(p)
    wb_f = openpyxl.load_workbook(p, data_only=False, read_only=False)
    wb_v = openpyxl.load_workbook(p, data_only=True, read_only=False)
    # arquivo salvo sem valores calculados (gerado por sistema): recalcula pelo LibreOffice
    formulas = [(ws.title, c.coordinate) for ws in wb_f.worksheets for row in ws.iter_rows() for c in row
                if isinstance(c.value, str) and c.value.startswith("=")]
    if formulas and sum(1 for t, k in formulas[:500] if wb_v[t][k].value is None) > 0.5 * min(len(formulas), 500):
        wb_v = openpyxl.load_workbook(converter(p), data_only=True, read_only=False)
    partes, total = [], 0
    for ws_f in wb_f.worksheets:
        ws_v = wb_v[ws_f.title]
        if ws_f.sheet_state != "visible":
            estado = " (aba oculta)"
        else:
            estado = ""
        linhas = []
        for row in ws_f.iter_rows():
            celulas = []
            for c in row:
                if c.value is None:
                    continue
                v = ws_v[c.coordinate].value
                col = get_column_letter(c.column)
                if isinstance(c.value, str) and c.value.startswith("="):
                    calc = fmt(v, c.number_format)
                    celulas.append(f"{col}={calc or '?'} [{c.value[:120]}]")
                else:
                    t = fmt(c.value, c.number_format)
                    if t:
                        celulas.append(f"{col}={t[:400]}")
            if celulas:
                linhas.append(f"L{row[0].row}: " + " | ".join(celulas))
        if not linhas:
            continue
        bloco = f"### Aba \"{ws_f.title}\"{estado} — {len(linhas)} linhas preenchidas\n" + "\n".join(linhas)
        if total + len(bloco) > limite:
            bloco = bloco[: max(0, limite - total)] + "\n[... planilha truncada por tamanho ...]"
            partes.append(bloco)
            break
        partes.append(bloco)
        total += len(bloco)
    return "\n\n".join(partes)


def main():
    if len(sys.argv) < 3:
        sys.exit(__doc__)
    limite = int(sys.argv[sys.argv.index("--max") + 1]) if "--max" in sys.argv else MAX
    texto = extrair(sys.argv[1], limite)
    if len(texto.strip()) < 50:
        print("planilha sem conteúdo legível", file=sys.stderr)
        sys.exit(3)
    pathlib.Path(sys.argv[2]).write_text(texto, encoding="utf-8")
    print(f"planilha extraída: {len(texto)} caracteres")


if __name__ == "__main__":
    main()
